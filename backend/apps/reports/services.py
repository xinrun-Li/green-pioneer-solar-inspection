import io
import logging
from datetime import datetime

from django.db.models import Count, Q, Value
from django.db.models.functions import Coalesce
from django.template.loader import render_to_string
from django.utils.timezone import localtime

from .models import Report

logger = logging.getLogger(__name__)


class ReportSnapshotService:
    """数据快照收集服务"""

    @classmethod
    def collect(cls, inspection_task=None):
        """从各模块收集统计数据"""
        station_id = getattr(inspection_task, "station_id", None)
        snapshot = {
            "total_uploads": 0,
            "total_detections": 0,
            "classification_counts": {"normal": 0, "needs_cleaning": 0, "needs_repair": 0},
            "total_abnormal_events": 0,
            "abnormal_events_by_status": {},
            "total_inspection_tasks": 0,
            "inspection_tasks_by_status": {},
            "station_coverage": {},
            "inspection_task": {
                "id": getattr(inspection_task, "id", None),
                "title": getattr(inspection_task, "title", ""),
                "status": getattr(inspection_task, "status", ""),
                "total_waypoints": getattr(inspection_task, "total_waypoints", 0),
                "visited_waypoints": getattr(inspection_task, "visited_waypoints", 0),
                "coverage": getattr(inspection_task, "coverage", 0),
            },
        }

        # 1. 上传统计
        try:
            from apps.uploads.models import MediaAsset

            uploads = MediaAsset.objects.all()
            if station_id:
                uploads = uploads.filter(batch__station_id=station_id)
            snapshot["total_uploads"] = uploads.count()
        except ImportError:
            logger.warning("apps.uploads 模块不可用，跳过上传统计")

        # 2. 识别统计
        try:
            from apps.recognition.models import Detection

            detections = Detection.objects.all()
            if station_id:
                detections = detections.filter(job__media__batch__station_id=station_id)
            snapshot["total_detections"] = detections.count()
            classification_counts = {}
            for row in (
                detections.filter(review_status=Detection.ReviewStatus.CONFIRMED).annotate(
                    effective=Coalesce("confirmed_class", "original_class")
                )
                .values("effective")
                .annotate(count=Count("id"))
            ):
                classification_counts[row["effective"]] = row["count"]
            snapshot["classification_counts"] = {
                "normal": classification_counts.get("normal", 0),
                "needs_cleaning": classification_counts.get("cleaning", 0),
                "needs_repair": classification_counts.get("repair", 0),
            }
        except ImportError:
            logger.warning("apps.recognition 模块不可用，跳过识别统计")

        # 3. 异常事件统计
        try:
            from apps.events.models import AbnormalEvent

            events = AbnormalEvent.objects.all()
            if station_id:
                events = events.filter(panel__array__region__station_id=station_id)
            snapshot["total_abnormal_events"] = events.count()
            status_counts = dict(
                events.values_list("status").annotate(count=Count("id"))
            )
            snapshot["abnormal_events_by_status"] = {
                "open": status_counts.get("open", 0),
                "closed": status_counts.get("closed", 0),
            }
        except ImportError:
            logger.warning("apps.events 模块不可用，跳过异常事件统计")

        # 4. 巡检任务统计
        try:
            from apps.inspections.models import InspectionTask

            tasks = InspectionTask.objects.filter(id=inspection_task.id) if inspection_task else InspectionTask.objects.all()
            if station_id:
                tasks = tasks.filter(station_id=station_id)
            snapshot["total_inspection_tasks"] = tasks.count()
            status_counts = dict(
                tasks.values_list("status").annotate(count=Count("id"))
            )
            snapshot["inspection_tasks_by_status"] = status_counts
        except ImportError:
            logger.warning("apps.inspections 模块不可用，跳过巡检任务统计")

        # 5. 电站覆盖统计
        try:
            from apps.stations.models import Panel, Region, SolarArray

            panels = Panel.objects.all()
            arrays = SolarArray.objects.all()
            regions = Region.objects.all()
            if station_id:
                panels = panels.filter(array__region__station_id=station_id)
                arrays = arrays.filter(region__station_id=station_id)
                regions = regions.filter(station_id=station_id)
            snapshot["station_coverage"] = {
                "total_panels": panels.count(),
                "total_arrays": arrays.count(),
                "total_regions": regions.count(),
                "regions": list(
                    regions.annotate(
                        array_count=Count("arrays", distinct=True),
                        panel_count=Count("arrays__panels", distinct=True),
                    ).values("id", "name", "direction", "array_count", "panel_count")
                ),
            }
        except ImportError:
            logger.warning("apps.stations 模块不可用，跳过电站覆盖统计")

        return snapshot


class ReportGenerator:
    """报告生成器"""

    @classmethod
    def generate(cls, report):
        """根据报告类型路由到对应生成器"""
        report.progress = 5
        report.save(update_fields=["progress"])

        try:
            # 先收集快照数据
            if not report.snapshot_data:
                report.snapshot_data = ReportSnapshotService.collect()
                report.save(update_fields=["snapshot_data"])

            report.progress = 15
            report.save(update_fields=["progress"])

            generators = {
                Report.ReportType.WEB: cls._generate_web,
                Report.ReportType.EXCEL: cls._generate_excel,
                Report.ReportType.PDF: cls._generate_pdf,
            }
            generator = generators.get(report.report_type)
            if generator is None:
                raise ValueError(f"不支持的报告类型: {report.report_type}")

            generator(report)

            report.progress = 100
            report.status = Report.Status.READY
            report.save(update_fields=["progress", "status"])
        except Exception as e:
            logger.exception("报告生成失败")
            report.status = Report.Status.FAILED
            report.error_message = str(e)
            report.save(update_fields=["status", "error_message"])

    @classmethod
    def _generate_web(cls, report):
        """生成 HTML 网页报告"""
        snapshot_data = report.snapshot_data
        context = {
            "report": report,
            "snapshot_data": snapshot_data,
            "generated_at": localtime(report.created_at).strftime("%Y-%m-%d %H:%M:%S"),
        }
        html_content = render_to_string("reports/report_web.html", context)

        # 保存为 HTML 文件
        from django.core.files.base import ContentFile

        report.file.save(
            f"report_{report.id}.html",
            ContentFile(html_content.encode("utf-8")),
        )

        report.progress = 90
        report.save(update_fields=["progress"])

    @classmethod
    def _generate_excel(cls, report):
        """生成 Excel 报告"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        snapshot_data = report.snapshot_data

        # Sheet 1: 概览
        ws1 = wb.active
        ws1.title = "汇总统计"
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=12, color="FFFFFF")

        ws1.cell(row=1, column=1, value="指标").font = header_font_white
        ws1.cell(row=1, column=1).fill = header_fill
        ws1.cell(row=1, column=2, value="数值").font = header_font_white
        ws1.cell(row=1, column=2).fill = header_fill

        metrics = [
            ("总上传数", snapshot_data.get("total_uploads", 0)),
            ("总识别数", snapshot_data.get("total_detections", 0)),
            ("正常组件", snapshot_data.get("classification_counts", {}).get("normal", 0)),
            ("需要清洗", snapshot_data.get("classification_counts", {}).get("needs_cleaning", 0)),
            ("需要维修", snapshot_data.get("classification_counts", {}).get("needs_repair", 0)),
            ("异常事件数", snapshot_data.get("total_abnormal_events", 0)),
            ("巡检任务数", snapshot_data.get("total_inspection_tasks", 0)),
        ]
        for i, (key, value) in enumerate(metrics, start=2):
            ws1.cell(row=i, column=1, value=key)
            ws1.cell(row=i, column=2, value=value)

        ws1.column_dimensions["A"].width = 20
        ws1.column_dimensions["B"].width = 15

        # Sheet 2: 异常明细
        ws2 = wb.create_sheet("组件明细")
        ws2.cell(row=1, column=1, value="ID").font = header_font_white
        ws2.cell(row=1, column=1).fill = header_fill
        ws2.cell(row=1, column=2, value="组件").font = header_font_white
        ws2.cell(row=1, column=2).fill = header_fill
        ws2.cell(row=1, column=3, value="类型").font = header_font_white
        ws2.cell(row=1, column=3).fill = header_fill
        ws2.cell(row=1, column=4, value="状态").font = header_font_white
        ws2.cell(row=1, column=4).fill = header_fill
        ws2.cell(row=1, column=5, value="开启时间").font = header_font_white
        ws2.cell(row=1, column=5).fill = header_fill

        try:
            from apps.events.models import AbnormalEvent

            events = AbnormalEvent.objects.select_related("panel__array__region").order_by("-opened_at")
            for i, event in enumerate(events, start=2):
                ws2.cell(row=i, column=1, value=event.id)
                ws2.cell(row=i, column=2, value=event.panel.full_code)
                ws2.cell(row=i, column=3, value=event.get_event_type_display())
                ws2.cell(row=i, column=4, value=event.get_status_display())
                ws2.cell(row=i, column=5, value=localtime(event.opened_at).strftime("%Y-%m-%d %H:%M"))
        except ImportError:
            ws2.cell(row=2, column=1, value="异常事件模块不可用")

        ws2.column_dimensions["A"].width = 10
        ws2.column_dimensions["B"].width = 25
        ws2.column_dimensions["C"].width = 12
        ws2.column_dimensions["D"].width = 10
        ws2.column_dimensions["E"].width = 20

        # Sheet 3: 电站覆盖
        ws3 = wb.create_sheet("区域汇总")
        ws3.cell(row=1, column=1, value="区域").font = header_font_white
        ws3.cell(row=1, column=1).fill = header_fill
        ws3.cell(row=1, column=2, value="方向").font = header_font_white
        ws3.cell(row=1, column=2).fill = header_fill
        ws3.cell(row=1, column=3, value="阵列数").font = header_font_white
        ws3.cell(row=1, column=3).fill = header_fill
        ws3.cell(row=1, column=4, value="组件数").font = header_font_white
        ws3.cell(row=1, column=4).fill = header_fill

        regions = snapshot_data.get("station_coverage", {}).get("regions", [])
        for i, region in enumerate(regions, start=2):
            ws3.cell(row=i, column=1, value=region.get("name", ""))
            ws3.cell(row=i, column=2, value=region.get("direction", ""))
            ws3.cell(row=i, column=3, value=region.get("array_count", 0))
            ws3.cell(row=i, column=4, value=region.get("panel_count", 0))

        ws3.column_dimensions["A"].width = 20
        ws3.column_dimensions["B"].width = 12
        ws3.column_dimensions["C"].width = 10
        ws3.column_dimensions["D"].width = 10

        # 保存
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        from django.core.files.base import ContentFile

        report.file.save(
            f"report_{report.id}.xlsx",
            ContentFile(output.read()),
        )

        report.progress = 90
        report.save(update_fields=["progress"])

    @classmethod
    def _generate_pdf(cls, report):
        """生成 PDF 报告"""
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        snapshot_data = report.snapshot_data
        buffer = io.BytesIO()

        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            title=f"报告 #{report.id}",
            topMargin=20 * mm,
            bottomMargin=20 * mm,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            fontSize=18,
            spaceAfter=12,
        )
        subtitle_style = ParagraphStyle(
            "Subtitle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.grey,
            spaceAfter=20,
        )
        heading_style = ParagraphStyle(
            "SectionHeading",
            parent=styles["Heading2"],
            fontSize=14,
            spaceBefore=16,
            spaceAfter=8,
        )

        elements = []
        elements.append(Paragraph(f"绿能先锋 · 光伏电站智能识别系统报告", title_style))
        elements.append(
            Paragraph(
                f"生成时间: {localtime(report.created_at).strftime('%Y-%m-%d %H:%M:%S')}",
                subtitle_style,
            )
        )
        elements.append(Spacer(1, 10 * mm))

        # 概览表格
        elements.append(Paragraph("概览数据", heading_style))
        overview_data = [
            ["指标", "数值"],
            ["总上传数", str(snapshot_data.get("total_uploads", 0))],
            ["总识别数", str(snapshot_data.get("total_detections", 0))],
            [
                "正常组件",
                str(snapshot_data.get("classification_counts", {}).get("normal", 0)),
            ],
            [
                "需要清洗",
                str(snapshot_data.get("classification_counts", {}).get("needs_cleaning", 0)),
            ],
            [
                "需要维修",
                str(snapshot_data.get("classification_counts", {}).get("needs_repair", 0)),
            ],
            ["异常事件数", str(snapshot_data.get("total_abnormal_events", 0))],
            ["巡检任务数", str(snapshot_data.get("total_inspection_tasks", 0))],
        ]
        overview_table = Table(overview_data, colWidths=[120, 80])
        overview_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 10),
                    ("FONTSIZE", (0, 1), (-1, -1), 9),
                    ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ]
            )
        )
        elements.append(overview_table)

        # 构建 PDF
        doc.build(elements)
        buffer.seek(0)

        from django.core.files.base import ContentFile

        report.file.save(
            f"report_{report.id}.pdf",
            ContentFile(buffer.read()),
        )

        report.progress = 90
        report.save(update_fields=["progress"])


class ReportRetryService:
    """报告重试服务"""

    @classmethod
    def retry(cls, report):
        """重置报告状态并重新生成"""
        report.status = Report.Status.GENERATING
        report.progress = 0
        report.error_message = ""
        report.file.delete(save=False)
        report.save(update_fields=["status", "progress", "error_message", "file"])

        ReportGenerator.generate(report)
