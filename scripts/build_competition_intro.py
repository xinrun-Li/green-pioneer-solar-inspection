from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/lixinrun/Documents/绿能先锋/绿能先锋_参赛作品简介.docx")

FONT = "Hiragino Sans GB"
FONT_LATIN = "Aptos"
INK = "183B56"
TEAL = "0F766E"
GREEN = "15803D"
ORANGE = "C2410C"
MUTED = "52606D"
LIGHT = "E8F5F1"
LIGHT_BLUE = "EEF5F8"
LIGHT_ORANGE = "FFF4E8"
RULE = "D5E1E2"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs.get(edge)
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name=FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para(paragraph, before=0, after=6, line=1.25, align=None, keep=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if keep:
        fmt.keep_with_next = True


def add_text(doc, text, size=10.5, color=INK, bold=False, after=7, align=None, italic=False, keep=False):
    p = doc.add_paragraph()
    set_para(p, after=after, align=align, keep=keep)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_section_heading(doc, title: str, eyebrow: str | None = None):
    if eyebrow:
        p = doc.add_paragraph()
        set_para(p, before=14, after=2, line=1.0, keep=True)
        r = p.add_run(eyebrow.upper())
        set_run_font(r, size=8.5, color=TEAL, bold=True)
    p = doc.add_paragraph()
    set_para(p, before=0, after=6, line=1.0, keep=True)
    r = p.add_run(title)
    set_run_font(r, size=15, color=INK, bold=True)
    return p


def add_feature(doc, label: str, body: str, accent=TEAL):
    p = doc.add_paragraph()
    set_para(p, before=4, after=3, line=1.15, keep=True)
    tag = p.add_run(label)
    set_run_font(tag, size=10.5, color=accent, bold=True)
    desc = p.add_run("  " + body)
    set_run_font(desc, size=10.5, color=INK)
    return p


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, before=4, after=0, line=1.0)
    r = p.add_run("绿能先锋  |  参赛作品简介  ·  ")
    set_run_font(r, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)
    set_run_font(r, size=8.5, color=MUTED)


def add_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para(p, before=0, after=0, line=1.0)
    left = p.add_run("GREEN ENERGY OPERATIONS")
    set_run_font(left, size=8, color=TEAL, bold=True)
    right = p.add_run("  /  LOCAL AI INSPECTION PLATFORM")
    set_run_font(right, size=8, color=MUTED)


def add_metadata_table(doc):
    rows = [
        ("作品名称", "绿能先锋", "作品方向", "绿色能源 / 智能运维"),
        ("服务对象", "光伏电站控制人员", "核心形态", "本地无人机巡检分析系统"),
        ("关键词", "AI 视觉识别 · 人机协同 · 闭环追溯", "当前版本", "V1 业务闭环与 Mock 适配器"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    set_table_geometry(table, [1250, 3380, 1250, 3480], indent_dxa=120)
    mark_header_row(table.rows[0])
    for row, values in zip(table.rows, rows):
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            set_para(p, after=0, line=1.1)
            r = p.add_run(value)
            set_run_font(r, size=9 if i % 2 == 0 else 9.2, color=TEAL if i % 2 == 0 else INK, bold=i % 2 == 0)
            if i % 2 == 0:
                set_cell_shading(cell, LIGHT)
            else:
                set_cell_shading(cell, WHITE)
            border = {"val": "single", "sz": 4, "color": RULE, "space": 0}
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title: str, body: str, fill=LIGHT_BLUE, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    border = {"val": "single", "sz": 7, "color": RULE, "space": 0}
    set_cell_border(cell, top=border, bottom=border, left={"val": "single", "sz": 18, "color": accent, "space": 0}, right=border)
    cell.text = ""
    p = cell.paragraphs[0]
    set_para(p, after=4, line=1.18, keep=True)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=accent, bold=True)
    p2 = cell.add_paragraph()
    set_para(p2, after=0, line=1.32)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    add_header(section)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    # Opening block: proposal_centerpiece style adapted for a concise competition brief.
    add_text(doc, "参赛作品简介", size=9.5, color=TEAL, bold=True, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "绿能先锋", size=28, color=INK, bold=True, after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "面向光伏电站控制人员的本地无人机巡检分析系统", size=13.5, color=MUTED, after=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(doc, "让巡检影像从“看见异常”走向“确认、处置与追溯”", size=11, color=ORANGE, bold=True, after=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_metadata_table(doc)

    add_callout(
        doc,
        "一段话摘要",
        "“绿能先锋”聚焦光伏电站无人机巡检中影像分散、异常判断依赖经验、任务执行与结果留痕脱节等问题，构建集巡检影像上传、AI 视觉初筛、低置信度人工复核、组件地图、异常事件、巡检任务、报告导出和历史分析于一体的本地化运维平台。系统以 PostgreSQL 作为业务事实来源，以本地规则引擎和推理适配器支撑断网条件下的连续工作，并对所有写入动作设置人工确认门槛，形成可解释、可追溯、可复盘的光伏运维闭环。",
        fill=LIGHT,
        accent=GREEN,
    )

    add_section_heading(doc, "一、作品定位与应用场景", "Positioning")
    add_text(
        doc,
        "光伏电站规模扩大后，巡检人员需要同时处理无人机图片或视频、组件状态、异常事件、维修清洗任务和阶段性报告。传统流程往往停留在“拍摄后人工查看”，识别结果难以直接对应到具体组件，异常也容易在复核、处置和统计之间断裂。绿能先锋从控制人员的实际工作台出发，将一次巡检拆解为可确认、可执行、可回溯的业务步骤：先上传并识别，再由人复核，随后联动地图与异常事件，最终沉淀为任务进度、报告和历史趋势。",
        after=7,
    )
    add_text(
        doc,
        "作品适用于集中式或分布式光伏电站的日常巡检、异常筛查、清洗维修协同和运维汇报，也可作为后续接入真实 YOLO 模型、飞控数据和红外复检能力的业务底座。",
        after=7,
    )

    add_section_heading(doc, "二、核心功能与工作闭环", "Core workflow")
    add_feature(doc, "影像识别与人工复核", "支持 JPG、PNG、WEBP 图片及 MP4、MOV 视频上传，输出“正常、需要清洗、需要维修”三类结果；低于 0.75 置信度的结果进入复核，确认前不写入最终地图状态、报告或统计。", TEAL)
    add_feature(doc, "组件地图与异常事件", "以区域、阵列和组件为组织单元，展示 4 个区域、12 个阵列、240 块组件的状态网格；点击组件即可查看当前位置、异常事件与历史记录，并支持异常开启、关闭和再次异常留痕。", GREEN)
    add_feature(doc, "巡检任务与航点执行", "通过“草稿 - 确认 - 执行 - 暂停/恢复 - 完成”的状态机管理任务，生成 S 形巡检路线和连续航点，记录任务进度、覆盖率、素材关联、跳过原因与执行时间线。", ORANGE)
    add_feature(doc, "本地助手与报告沉淀", "内置 10 类本地意图识别能力，可查询状态、异常、任务和趋势，也可生成任务或报告草稿；所有写操作须人工确认。系统同时生成 HTML、Excel、PDF 报告，并保留历史统计与异常开闭记录。", TEAL)

    add_section_heading(doc, "三、作品创新与特色", "What is different")
    add_feature(doc, "从识别工具到运维闭环", "不把 AI 识别作为孤立功能，而是把识别结果接入组件地图、异常事件、巡检任务和报告中心，使一次检测能够进入后续处置流程。", TEAL)
    add_feature(doc, "本地优先的人机协同", "本地规则引擎优先承担固定业务查询和统计，云端能力只作为非预设问题的回退；助手生成的任务、关闭事件、导出报告等动作均先形成草稿，控制人员确认后才执行。", GREEN)
    add_feature(doc, "面向事实来源的可追溯设计", "以 PostgreSQL 保存业务事实，Redis 只承担队列和临时进度；识别结果、人工改判、事件变化、任务状态和报告快照均保留关联关系，便于复核、审计和问题定位。", ORANGE)
    add_feature(doc, "可替换的模型接入架构", "通过统一推理适配器隔离模型与业务 API，当前 V1 可使用明确标记的 Mock 适配器完成闭环验证，后续可平滑接入 YOLOv11-nano 等真实视觉模型而不改变既有业务接口。", TEAL)

    add_section_heading(doc, "四、技术实现与安全边界", "Implementation")
    add_text(
        doc,
        "系统采用 Vue 3 + TypeScript + Vite 构建前端控制台，使用 Django 5.2 + Django REST Framework 提供业务 API，以 PostgreSQL 管理电站、组件、识别、事件、任务和报告等持久化数据，Redis + Celery 支撑异步任务，并预留 macOS 原生 AI Worker 进行本地推理。认证、角色、电站归属、状态机、文件校验、去重和失败重试均由后端约束，前端不直接绕过业务规则。",
        after=7,
    )
    add_callout(
        doc,
        "当前 V1 的真实边界",
        "当前版本的默认验收重点是业务闭环、权限门禁、人工确认、状态机和统计一致性；真实 YOLO 视频推理、完整训练执行器、飞控接入和红外诊断属于后续专项能力。对于 RGB 影像中的黑化现象，系统统一表述为“疑似热斑，建议红外复检”，不将其当作确诊结论。",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    add_section_heading(doc, "五、应用价值与参赛意义", "Impact")
    add_text(
        doc,
        "对一线控制人员，绿能先锋把分散的影像、地图、任务和报告集中到同一工作台，减少重复查找和口径不一致；对电站管理者，系统让清洗维修对象、任务进度和异常处理过程具备结构化记录，便于安排资源和复盘运营；对后续智能化升级，平台通过适配器、困难样本、数据集版本和模型版本接口，为真实模型迭代和规模化部署保留扩展空间。作品的价值不只在于“能识别”，更在于让识别结果能够被人确认、被系统执行，并在之后被查证。",
        after=5,
    )
    add_text(doc, "关键词：光伏运维｜无人机巡检｜AI 视觉识别｜人工复核｜本地化部署｜异常闭环｜可追溯管理", size=9.5, color=TEAL, bold=True, after=0)

    # Keep core properties neutral for external submission.
    doc.core_properties.title = "绿能先锋 - 参赛作品简介"
    doc.core_properties.subject = "光伏电站本地无人机巡检分析系统"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
